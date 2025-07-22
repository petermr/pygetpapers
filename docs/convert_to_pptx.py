#!/usr/bin/env python3
"""
Convert Markdown presentation to PowerPoint format.
"""

import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


def parse_markdown_slides(markdown_file):
    """Parse markdown file and extract slides."""
    with open(markdown_file, "r", encoding="utf-8") as f:
        content = f.read()

    # Split by slide separators (---)
    slides = re.split(r"\n---\n", content)

    parsed_slides = []
    for slide in slides:
        if slide.strip():
            # Extract title and content
            lines = slide.strip().split("\n")
            title = ""
            content = []

            for line in lines:
                if line.startswith("# ") and not title:
                    title = line[2:].strip()
                elif line.startswith("## "):
                    # Subtitle
                    content.append(("subtitle", line[3:].strip()))
                elif line.startswith("### "):
                    # Subheading
                    content.append(("subheading", line[4:].strip()))
                elif line.startswith("**") and line.endswith("**"):
                    # Bold text
                    content.append(("bold", line[2:-2].strip()))
                elif line.startswith("- "):
                    # Bullet point
                    content.append(("bullet", line[2:].strip()))
                elif line.startswith("```"):
                    # Code block
                    content.append(("code", line[3:].strip()))
                elif line.strip() and not line.startswith("#"):
                    # Regular text
                    content.append(("text", line.strip()))

            parsed_slides.append({"title": title, "content": content})

    return parsed_slides


def create_presentation(slides_data, output_file):
    """Create PowerPoint presentation from parsed slides."""
    prs = Presentation()

    # Set slide dimensions (16:9 aspect ratio)
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)

    for slide_data in slides_data:
        # Create slide with title and content layout
        slide_layout = prs.slide_layouts[1]  # Title and Content layout
        slide = prs.slides.add_slide(slide_layout)

        # Set title
        title = slide.shapes.title
        title.text = slide_data["title"]
        title.text_frame.paragraphs[0].font.size = Pt(44)
        title.text_frame.paragraphs[0].font.bold = True
        title.text_frame.paragraphs[0].font.color.rgb = RGBColor(31, 73, 125)

        # Set content
        content_placeholder = slide.placeholders[1]
        text_frame = content_placeholder.text_frame
        text_frame.clear()

        for item_type, text in slide_data["content"]:
            if item_type == "subtitle":
                p = text_frame.add_paragraph()
                p.text = text
                p.font.size = Pt(32)
                p.font.bold = True
                p.font.color.rgb = RGBColor(68, 84, 106)
                p.space_after = Pt(12)

            elif item_type == "subheading":
                p = text_frame.add_paragraph()
                p.text = text
                p.font.size = Pt(24)
                p.font.bold = True
                p.font.color.rgb = RGBColor(89, 89, 89)
                p.space_after = Pt(6)

            elif item_type == "bold":
                p = text_frame.add_paragraph()
                p.text = text
                p.font.size = Pt(18)
                p.font.bold = True
                p.font.color.rgb = RGBColor(0, 0, 0)
                p.space_after = Pt(6)

            elif item_type == "bullet":
                p = text_frame.add_paragraph()
                p.text = f"• {text}"
                p.font.size = Pt(16)
                p.font.color.rgb = RGBColor(0, 0, 0)
                p.level = 0
                p.space_after = Pt(3)

            elif item_type == "code":
                p = text_frame.add_paragraph()
                p.text = text
                p.font.size = Pt(14)
                p.font.name = "Courier New"
                p.font.color.rgb = RGBColor(0, 0, 139)
                p.space_after = Pt(6)

            elif item_type == "text":
                p = text_frame.add_paragraph()
                p.text = text
                p.font.size = Pt(16)
                p.font.color.rgb = RGBColor(0, 0, 0)
                p.space_after = Pt(3)

    # Save presentation
    prs.save(output_file)
    print(f"✅ PowerPoint presentation saved: {output_file}")


def main():
    """Main function to convert markdown to PowerPoint."""
    # File paths
    markdown_file = Path(__file__).parent / "upspace-implementation-presentation.md"
    output_file = Path(__file__).parent / "UPSpace_Implementation_Presentation.pptx"

    print("=" * 60)
    print("Converting Markdown to PowerPoint")
    print("=" * 60)

    if not markdown_file.exists():
        print(f"❌ Markdown file not found: {markdown_file}")
        return

    try:
        # Parse markdown slides
        print("📖 Parsing markdown slides...")
        slides_data = parse_markdown_slides(markdown_file)
        print(f"✅ Found {len(slides_data)} slides")

        # Create PowerPoint presentation
        print("📊 Creating PowerPoint presentation...")
        create_presentation(slides_data, output_file)

        print(f"\n🎉 Conversion complete!")
        print(f"📁 Output file: {output_file.absolute()}")
        print(f"📊 Slides created: {len(slides_data)}")

    except ImportError:
        print("❌ python-pptx library not installed.")
        print("💡 Install it with: pip install python-pptx")
    except Exception as e:
        print(f"❌ Error during conversion: {e}")


if __name__ == "__main__":
    main()
